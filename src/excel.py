# Writes several tables to one Excel workbook with a readme sheet.
# Used by: render.py.
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill('solid', start_color='DDE4EE')
MAX_ROWS = 1_000_000


def style_sheet(ws, ncols, widths=None):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name='Arial', bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 14)
    ws.freeze_panes = 'A2'


def set_default_font(wb, name='Arial'):
    for st in wb._named_styles:
        if st.name == 'Normal':
            st.font = Font(name=name, size=10)


def write_workbook(path, sheets, readme=None):
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        if readme:
            pd.DataFrame(readme, columns=['sheet', 'what it contains']).to_excel(xw, sheet_name='README', index=False)
        for name, df in sheets.items():
            df = df.iloc[:MAX_ROWS]
            df.to_excel(xw, sheet_name=name[:31], index=False)
        wb = xw.book
        set_default_font(wb)
        for ws in wb.worksheets:
            style_sheet(ws, ws.max_column, {2: 110} if ws.title == 'README' else None)
            if ws.title == 'README':
                for row in ws.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical='top')
